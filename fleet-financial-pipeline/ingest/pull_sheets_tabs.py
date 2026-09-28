"""Pull individual, recent tabs from a P&L Google Sheet via the Sheets API.

WHY THIS EXISTS, NOT pull_sheets.py'S EXPORT. Drive's files.export endpoint
renders the ENTIRE spreadsheet before handing back .xlsx, and Google enforces
a hard size ceiling on that render -- ZONE (139 tabs) and XTRACK (141 tabs)
both exceed it regardless of which credential calls it. Confirmed live
2026-09-28: the Drive connector AND this project's own service account both
hit "This file is too large to be exported" on the identical files -- the
service account changed nothing about this limit, because it is not a
permissions problem. A single tab is a few KB; the Sheets API's
spreadsheets.values.get reads one tab at a time and never renders the rest of
the spreadsheet, so the size ceiling never applies.

WHAT THIS DOES. For one company, lists every tab in its sheet, keeps the ones
ingest_weekly_pnl.week_key() can date, and fetches -- via values.get,
valueRenderOption=UNFORMATTED_VALUE, i.e. the same computed numbers
data_only=True already reads out of a full export, never a formula string --
every tab whose week is new: strictly after the latest week already in the
local workbook, plus that latest week itself (in case that tab was still
being edited when the last export ran).

HOW IT MERGES WITHOUT RISKING THE EXISTING FILE. The local .xlsx the pipeline
already reads (ingest_weekly_pnl.WORKBOOKS[company]) contains real Excel
FORMULAS with cached values -- openpyxl only ever exposes the cached values
(data_only=True), and if this script re-opened that file in write mode,
openpyxl would drop those cached values on save (a documented openpyxl
limitation: it does not recompute formulas, and does not preserve the <v>
cache for a formula cell through a load/save round trip). That would silently
zero every EXISTING week the next time anything reads the file with
data_only=True -- the same access pattern every consumer in this pipeline
already uses. So this never re-opens the existing file in write mode. Instead
it reads every existing tab's cell grid as already-computed VALUES
(data_only=True) and builds a brand-new workbook from those values plus the
newly-fetched tabs, all as plain values with no formulas at all. Verified by
round-tripping ZONE's real 26-week file exactly this way and confirming
ingest_weekly_pnl.read_workbook() parses the rebuilt file identically to the
original -- 0 diffs across all 26 weeks (2026-09-28). The new file is written
atomically (tmp + replace), so a run interrupted mid-write cannot leave a
partial workbook where the pipeline expects a whole one.

Needs `https://www.googleapis.com/auth/spreadsheets.readonly` in addition to
pull_sheets.py's drive.readonly. Both are read-only; this cannot write to the
source sheet, and does not try to.

    python3 ingest/pull_sheets_tabs.py ZONE_3YR             # fetch new tabs
    python3 ingest/pull_sheets_tabs.py XTRACK --check       # report only
    python3 ingest/pull_sheets_tabs.py ZONE_3YR --local-path /tmp/x.xlsx  # test target
"""
import argparse
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import pull_sheets as PS          # noqa: E402
import ingest_weekly_pnl as W     # noqa: E402

SCOPES = PS.SCOPES + ["https://www.googleapis.com/auth/spreadsheets.readonly"]

# pull_sheets.SHEETS keys -> the entity_id ingest_weekly_pnl.WORKBOOKS keys on.
COMPANY = {"ZONE_3YR": "ZONE", "XTRACK": "XTRACK", "AFG": "AFG"}


def sheets_service():
    from googleapiclient.discovery import build
    from google.oauth2 import service_account
    info = PS.read_key_info()
    creds = service_account.Credentials.from_service_account_info(info, scopes=SCOPES)
    return build("sheets", "v4", credentials=creds, cache_discovery=False)


def list_tabs(svc, spreadsheet_id):
    meta = svc.spreadsheets().get(
        spreadsheetId=spreadsheet_id, fields="sheets.properties.title").execute()
    return [s["properties"]["title"] for s in meta["sheets"]]


def fetch_tab_grid(svc, spreadsheet_id, tab):
    """Every cell of one tab, as computed values -- never a formula string.

    A1:Z3000: measured against a real ZONE tab (91 unit blocks, max_row 1214,
    max_col 26 == Z) 2026-09-28. The per-unit block table, not the summary
    panel, drives the row count -- a company with more trucks or a longer
    history of them needs more rows, so this is 2.5x that measured extent, not
    a tight fit. THE CONTROL IN main() IS WHAT CATCHES A RANGE THAT IS STILL
    TOO SMALL: a truncated table undercounts unit_gross against the panel's
    Total gross exactly the way a genuinely wrong parse would, so a control
    failure here means check this constant first, not just distrust the data.
    """
    resp = svc.spreadsheets().values().get(
        spreadsheetId=spreadsheet_id,
        range=f"'{tab}'!A1:Z3000",
        valueRenderOption="UNFORMATTED_VALUE").execute()
    return resp.get("values", [])


def has_financials(ws):
    """Whether a fetched tab has closed out at least one unit block with
    real (nonzero) money in it yet -- takes an already-built worksheet, not
    a raw grid, so it can reuse ingest_weekly_pnl.blocks() exactly rather
    than re-deriving its logic.

    A 'Total' row in column B is NOT enough by itself: a NOT-YET-FILLED
    current week can already carry every unit's row, driver name and 'Total'
    marker rolled over from the prior week's layout, with every dollar and
    mile genuinely at 0.0 because nothing has been settled yet. Confirmed
    live 2026-09-28: ZONE's and XTRACK's newest tab at the time had 93 real
    'Total' rows (ingest_weekly_pnl.blocks() found every one of them) and
    thousands of non-empty cells, yet check_weekly_pnl() read the whole tab
    as panel=0/units=0 -- every block's gross and miles were exactly zero. A
    first version of this check only tested for the 'Total' marker's
    presence and wrongly accepted this tab. Every OTHER consumer in this
    pipeline finds a week by iterating sheetnames through week_key(), with NO
    equivalent check -- so a tab like that, once written, reads as a real
    week everywhere except the one control built to catch it, and inflates
    every week-count and rolling-window comparison in the pipeline (this
    broke test_truck_weeks.py and multiple test_truck_breakeven.py controls
    the same day). Never write such a tab in; let the next run re-fetch and
    re-check it once the office actually closes the week out.
    """
    return any(gross or miles for gross, miles in W.blocks(ws))


def rebuild(local_path, new_tabs):
    """Write local_path atomically: every existing tab NOT in new_tabs is
    carried over as plain values from the current file; every tab in
    new_tabs (name -> list of rows, each a list of cell values) replaces or
    adds a sheet built fresh from that grid -- EXCEPT a tab with no closed-out
    unit blocks yet (see has_financials()), which is dropped entirely rather
    than written as a phantom zero week. Returns the set of tab names
    actually written from new_tabs, so callers can tell a real add from a
    silently-skipped not-yet-filled tab."""
    local_path = Path(local_path)
    old = openpyxl.load_workbook(local_path, data_only=True) if local_path.exists() else None
    out = openpyxl.Workbook()
    out.remove(out.active)

    staged = {}
    for name, rows in new_tabs.items():
        ws = out.create_sheet(title=name)
        for r, row in enumerate(rows, start=1):
            for cidx, val in enumerate(row, start=1):
                if val != "":
                    ws.cell(row=r, column=cidx, value=val)
        staged[name] = ws
    written = {name for name, ws in staged.items() if has_financials(ws)}
    for name in set(staged) - written:
        del out[name]  # staged but rejected -- not yet filled in, drop it

    if old is not None:
        for name in old.sheetnames:
            if name in written:
                continue  # replaced below with the freshly-pulled grid
            # A fetched-but-not-yet-filled tab (name in new_tabs, not in
            # written) falls through to here and keeps whatever this file
            # already had for it -- the safe fallback if a re-fetch of an
            # already-closed cutoff week ever came back looking empty.
            ws_old, ws_new = old[name], out.create_sheet(title=name)
            for row in ws_old.iter_rows():
                for c in row:
                    if c.value is not None:
                        ws_new.cell(row=c.row, column=c.column, value=c.value)
    # `written` sheets are already staged in `out` from the loop above --
    # nothing more to add for them here.
    if not out.sheetnames:
        # Nothing to write at all: no prior file, and every fetched tab was
        # rejected by has_financials(). openpyxl refuses to save a workbook
        # with zero visible sheets, and there is nothing worth writing
        # anyway -- leave local_path exactly as it was (absent).
        return written
    local_path.parent.mkdir(parents=True, exist_ok=True)
    tmp = local_path.with_suffix(".tmp")
    out.save(tmp)
    tmp.replace(local_path)
    return written


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("key", choices=sorted(COMPANY), help="a pull_sheets.SHEETS key")
    ap.add_argument("--local-path", help="override the WORKBOOKS destination (testing)")
    ap.add_argument("--check", action="store_true", help="report only, write nothing")
    a = ap.parse_args()

    sheet = PS.SHEETS[a.key]
    company = COMPANY[a.key]
    local_path = Path(a.local_path) if a.local_path else PS.ROOT / W.WORKBOOKS[company]

    existing = W.read_workbook(local_path) if local_path.exists() else {}
    cutoff = max(existing) if existing else None
    print(f"  local file: {local_path}")
    print(f"  latest local week: {cutoff or '(none)'}")

    svc = sheets_service()
    tabs = list_tabs(svc, sheet["id"])
    dated = sorted((t for t in tabs if W.week_key(t)), key=W.week_key)
    to_fetch = [t for t in dated if cutoff is None or W.week_key(t) >= cutoff]
    print(f"  {len(dated)} dated tabs on the sheet, {len(to_fetch)} to fetch "
          f"(>= {cutoff or 'all'})")
    for t in to_fetch:
        print(f"    {W.week_key(t)}  {t}")

    if a.check or not to_fetch:
        print("\n  --check only; nothing written." if a.check else
              "\n  Nothing new; nothing written.")
        return

    new_tabs = {t: fetch_tab_grid(svc, sheet["id"], t) for t in to_fetch}
    written = rebuild(local_path, new_tabs)
    skipped = [t for t in to_fetch if t not in written]

    refreshed = W.read_workbook(local_path)
    # The weeks THIS RUN actually wrote -- by week_key of the WRITTEN tabs
    # only (has_financials() may have dropped some), not a before/after set
    # difference. set(refreshed) - set(existing) misses the cutoff week on
    # purpose: it already existed in `existing`, so a set difference hides it
    # even though its VALUE was just replaced -- exactly the week most likely
    # to need the control, since it's re-fetched because it might have been
    # mid-edit last time.
    touched = sorted({W.week_key(t) for t in written})
    print(f"\n  wrote {local_path}")
    print(f"  weeks now: {len(refreshed)} (was {len(existing)}), "
          f"latest {max(refreshed) if refreshed else '(none)'}")
    print(f"  fetched/updated weeks: {touched}")
    if skipped:
        print(f"  skipped (fetched but not yet filled in -- no closed-out "
              f"unit blocks): {[W.week_key(t) for t in skipped]}")

    if not touched:
        print("  nothing written this run (every fetched tab was not yet filled in).")
        return
    _, bad = W.check_weekly_pnl({wk: refreshed[wk] for wk in touched if wk in refreshed})
    if bad:
        print("\n  CONTROL FAILED on:")
        for wk, panel, units, diff in bad:
            print(f"    {wk}  panel={panel}  units={units}  diff={diff}")
        print("  Do not trust these weeks' numbers until this is understood.")
    else:
        print("  control passed: panel gross == unit-block gross on every new week.")


if __name__ == "__main__":
    main()
