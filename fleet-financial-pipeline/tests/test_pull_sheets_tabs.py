"""Controls on the per-tab Sheets reader (ingest/pull_sheets_tabs.py).

No network here: these guard the rebuild logic against the real live bug
found 2026-09-28 -- a not-yet-filled current week already carries every unit's
row and a real 'Total' marker in column B, rolled over from the prior week's
layout, with every dollar and mile genuinely at 0.0. A check that only tests
for the 'Total' marker's presence (a first version of this module's) wrongly
accepts such a tab, and once written it reads as a real week everywhere in
the pipeline that finds a week by tab name alone -- inflating every
week-count and rolling-window comparison built on ingest_weekly_pnl.WORKBOOKS.
"""
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT / "ingest"))

import pull_sheets_tabs as PST  # noqa: E402
import ingest_weekly_pnl as W   # noqa: E402


def _sheet(wb, title, rows):
    ws = wb.create_sheet(title=title)
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    return ws


REAL_WEEK_ROWS = [
    ["1234", "Active"],
    ["", "Total", 5000.0, 800.0],
]
ZERO_WEEK_ROWS = [
    ["1234", "Active"],
    ["", "Total", 0.0, 0.0],
    ["5678", "Active"],
    ["", "Total", 0.0, 0.0],
]


def test_has_financials_accepts_a_real_closed_out_week():
    wb = openpyxl.Workbook()
    ws = _sheet(wb, "real", REAL_WEEK_ROWS)
    assert PST.has_financials(ws)


def test_has_financials_rejects_a_week_whose_blocks_are_all_zero():
    """The exact live bug: a 'Total' marker is present, but every block reads
    (0.0, 0.0) because nothing has settled yet for an in-progress week."""
    wb = openpyxl.Workbook()
    ws = _sheet(wb, "not yet filled", ZERO_WEEK_ROWS)
    assert W.blocks(ws), "the fixture must actually exercise a 'Total' row"
    assert not PST.has_financials(ws)


def test_has_financials_rejects_a_tab_with_no_total_row_at_all():
    wb = openpyxl.Workbook()
    ws = _sheet(wb, "empty", [["", ""], ["", ""]])
    assert not PST.has_financials(ws)


def test_rebuild_never_writes_a_not_yet_filled_tab(tmp_path):
    """The whole point: a brand-new week with zero real data must not appear
    in the rebuilt file at all -- not as an empty sheet, not as anything."""
    target = tmp_path / "company.xlsx"
    new_tabs = {"09.21.26-09.27.26 New": ZERO_WEEK_ROWS}
    written = PST.rebuild(target, new_tabs)

    assert written == set()
    assert not target.exists(), "nothing to write means nothing gets written"


def test_rebuild_keeps_a_real_new_week_and_carries_forward_the_rest(tmp_path):
    target = tmp_path / "company.xlsx"
    # Seed an existing file with one real week, the way pull_sheets_tabs.py
    # would find it before ever running.
    seed = openpyxl.Workbook()
    seed.remove(seed.active)
    _sheet(seed, "09.07.26-09.13.26", REAL_WEEK_ROWS)
    seed.save(target)

    new_tabs = {
        "09.14.26-09.20.26": REAL_WEEK_ROWS,
        "09.21.26-09.27.26 New": ZERO_WEEK_ROWS,
    }
    written = PST.rebuild(target, new_tabs)

    assert written == {"09.14.26-09.20.26"}
    weeks = W.read_workbook(target)
    assert set(weeks) == {"2026-09-07", "2026-09-14"}
    assert "2026-09-21" not in weeks


def test_rebuild_falls_back_to_the_old_tab_if_a_refetch_comes_back_empty(tmp_path):
    """A cutoff week is always re-fetched (it might have been mid-edit) -- if
    that refetch somehow comes back looking empty, the file must keep the
    LAST GOOD version of that week rather than silently losing it."""
    target = tmp_path / "company.xlsx"
    seed = openpyxl.Workbook()
    seed.remove(seed.active)
    _sheet(seed, "09.14.26-09.20.26", REAL_WEEK_ROWS)
    seed.save(target)

    new_tabs = {"09.14.26-09.20.26": ZERO_WEEK_ROWS}
    written = PST.rebuild(target, new_tabs)

    assert written == set()
    weeks = W.read_workbook(target)
    assert weeks["2026-09-14"]["unit_gross"] == pytest.approx(5000.0), (
        "the real old week must survive a refetch that came back empty")
